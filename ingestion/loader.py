import re
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import docx
from bs4 import BeautifulSoup
from pypdf import PdfReader

from ingestion.manifest import ManifestEntry


@dataclass(frozen=True)
class Section:
    """A run of body text under one heading, with the full heading path above it.

    Chunking works on these rather than on the flat text so a chunk boundary
    lands between provisions instead of in the middle of one.
    """

    heading_path: tuple[str, ...]
    text: str


@dataclass(frozen=True)
class Document:
    text: str
    source_doc: str
    effective_date: str
    program: str
    # Empty when a source exposes no headings at all; the chunker then falls
    # back to treating the whole document as one unlabelled section.
    sections: tuple[Section, ...] = ()


_DOCX_HEADING_STYLE = re.compile(r"Heading (\d)")


def _load_docx(path: Path) -> tuple[str, list[tuple[str, int]]]:
    document = docx.Document(str(path))
    lines: list[str] = []
    headings: list[tuple[str, int]] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name
        # The table of contents repeats every heading in the book alongside a
        # page number, which retrieves as a near-duplicate of the real section.
        if style.lower().startswith("toc"):
            continue
        lines.append(text)
        # Only an exact "Heading N" is a heading -- the SOP also uses styles like
        # "Heading 4 para under" for the body paragraph that follows one.
        match = _DOCX_HEADING_STYLE.fullmatch(style)
        if match:
            headings.append((text, int(match.group(1))))
    return "\n".join(lines), headings


_PAGE_NUMBER_LINE = re.compile(r"^page\s+\d+\s+of\s+\d+", re.IGNORECASE)
# How many lines at each edge of a page can be a running header or footer.
_PAGE_MARGIN_LINES = 3


def _strip_page_furniture(pages: list[list[str]]) -> list[str]:
    """Drop running headers and footers.

    They repeat on every page, so in the concatenated text they land in the
    middle of provisions and chunk into noise. Furniture is identified by
    position as well as by repetition: repetition alone would also match a
    definition or caveat that the document legitimately restates on several
    pages, so only lines sitting in a page's top or bottom margin qualify.
    """
    if len(pages) < 2:
        return [line for page in pages for line in page if line]

    margin_appearances: Counter[str] = Counter()
    for page in pages:
        margin = set(page[:_PAGE_MARGIN_LINES] + page[-_PAGE_MARGIN_LINES:])
        margin_appearances.update(line for line in margin if line)

    threshold = max(2, len(pages) // 2)
    furniture = {line for line, count in margin_appearances.items() if count >= threshold}

    return [
        line
        for page in pages
        for line in page
        if line and line not in furniture and not _PAGE_NUMBER_LINE.match(line)
    ]


_LIST_MARKERS = ("•", "◦", "-", "*", "–")


# A provision heading is a phrase; "TO:", "SUBJECT:" and "CONTROL NO.:" are
# document metadata and splitting them only yields a stub chunk.
_MIN_INLINE_HEADING_WORDS = 4


def _split_inline_headings(lines: list[str]) -> list[tuple[str, bool]]:
    """Break "Heading: first words of the body" into two lines.

    A PDF wraps text at a fixed width, so a heading short enough to leave room
    on its line runs on into its own body and is no longer a line ending in a
    colon. The provision then gets filed under the previous heading: in the fee
    notice the short-term fee rule lands under "For loans with a maturity that
    exceeds 12 months", labelling the chunk with the opposite of what it answers.

    List markers are excluded because a bullet such as "For loans of $150,000
    or less: 2% ..." is one row of a fee table, and promoting each row to a
    heading would break the table into unreadable fragments.

    Returns each line with whether it may still be considered for heading
    detection. The tail of a split is body text, and has to be marked: it
    begins mid-sentence on a capitalised word, so "The Annual Service Fee is
    set for" would otherwise read as a heading in its own right.
    """
    out: list[tuple[str, bool]] = []
    for line in lines:
        head, separator, rest = line.partition(": ")
        if (
            separator
            and rest.strip()
            and not line.startswith(_LIST_MARKERS)
            and len(head) <= 100
            and len(re.findall(r"\S+", head)) >= _MIN_INLINE_HEADING_WORDS
            # ". " means the prefix spans a sentence boundary, so this is prose
            # that happens to contain a colon rather than a heading.
            and ". " not in head
        ):
            out.append((f"{head}:", True))
            out.append((rest.strip(), False))
            continue
        out.append((line, True))
    return out


def _looks_like_heading(line: str) -> bool:
    """Heuristic heading test for PDFs, which carry no structural markup."""
    if not line or len(line) > 100:
        return False
    # "Upfront Fee for EWCP loans:" -- a trailing colon introduces a provision.
    if line.endswith(":"):
        return True
    # "Additional Upfront Fee for 7(a) Loan Increases" -- a short, mostly
    # capitalised line that is not a sentence.
    if len(line) > 80 or line.endswith((".", ";", ",")):
        return False
    words = re.findall(r"[A-Za-z]+", line)
    if not words:
        return False
    return sum(word[0].isupper() for word in words) / len(words) >= 0.5


def _load_pdf(path: Path) -> tuple[str, list[tuple[str, int]]]:
    reader = PdfReader(str(path))
    # Kept page by page: _strip_page_furniture needs to know where on a page a
    # line sat, which is lost once the pages are concatenated.
    pages = [
        [line.strip() for line in (page.extract_text() or "").split("\n")]
        for page in reader.pages
    ]
    marked = _split_inline_headings(_strip_page_furniture(pages))
    # A PDF has no heading levels to read, so every detected heading is a peer.
    headings = [
        (line, 1) for line, may_be_heading in marked if may_be_heading and _looks_like_heading(line)
    ]
    return "\n".join(line for line, _ in marked), headings


# Content containers to prefer, most specific first. eCFR wraps the actual
# regulation in div.part; taking the whole page instead pulls in ~50k chars of
# nav, cookie banners and "unsupported browser" notices, which chunk into dozens
# of near-duplicate fragments that dilute retrieval.
_HTML_CONTENT_SELECTORS = ("div.part", "main", "article", "#content")
_HTML_HEADING_TAGS = ("h1", "h2", "h3", "h4", "h5", "h6")


def _load_html(path: Path) -> tuple[str, list[tuple[str, int]]]:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    container = soup
    for selector in _HTML_CONTENT_SELECTORS:
        matches = soup.select(selector)
        if matches:
            # A page can repeat a selector for small fragments (eCFR has a stub
            # div.part alongside the real one); the real content is the largest.
            container = max(matches, key=lambda m: len(m.get_text()))
            break
    headings = [
        (heading.get_text(separator="\n").strip(), int(tag[1]))
        for heading in container.find_all(_HTML_HEADING_TAGS)
        for tag in [heading.name]
        if heading.get_text(strip=True)
    ]
    return container.get_text(separator="\n"), headings


def _build_sections(text: str, headings: list[tuple[str, int]]) -> tuple[Section, ...]:
    """Split the flat text at each heading, tracking the enclosing heading path.

    Headings are located by scanning forward, never by searching the whole
    string, so a heading whose wording repeats later in the document still
    anchors to its own occurrence.
    """
    spans: list[tuple[int, str, int]] = []
    cursor = 0
    for heading, level in headings:
        index = text.find(heading, cursor)
        if index == -1:
            continue
        spans.append((index, heading, level))
        cursor = index + len(heading)

    sections: list[Section] = []
    open_headings: dict[int, str] = {}
    for position, (index, heading, level) in enumerate(spans):
        end = spans[position + 1][0] if position + 1 < len(spans) else len(text)
        # A heading closes every deeper one still open above it.
        open_headings = {lvl: h for lvl, h in open_headings.items() if lvl < level}
        open_headings[level] = heading
        body = text[index + len(heading) : end].strip()
        if body:
            path = tuple(open_headings[lvl] for lvl in sorted(open_headings))
            sections.append(Section(heading_path=path, text=body))
    return tuple(sections)


_LOADERS = {"docx": _load_docx, "pdf": _load_pdf, "html": _load_html}


def load_document(entry: ManifestEntry, raw_dir: Path) -> Document:
    path = raw_dir / entry.filename
    text, headings = _LOADERS[entry.file_type](path)
    extracted = len(text.strip())
    if not extracted:
        raise ValueError(f"No text extracted from {path}")
    if extracted < entry.min_chars:
        raise ValueError(
            f"Only {extracted} chars extracted from {path} (expected at least "
            f"{entry.min_chars}) -- the download or parse likely failed"
        )
    return Document(
        text=text,
        source_doc=entry.filename,
        effective_date=entry.effective_date,
        program=entry.program,
        sections=_build_sections(text, headings),
    )
