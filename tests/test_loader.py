from unittest.mock import Mock, patch

import docx
from docx.enum.style import WD_STYLE_TYPE

from ingestion.loader import load_document
from ingestion.manifest import ManifestEntry


def test_load_docx_extracts_paragraph_text(tmp_path):
    document = docx.Document()
    document.add_paragraph("SBA loans require a minimum equity injection.")
    document.add_paragraph("")
    document.add_paragraph("See Chapter 3 for eligibility.")
    document.save(str(tmp_path / "sop.docx"))
    entry = ManifestEntry(
        filename="sop.docx", url="https://example.com/sop.docx",
        file_type="docx", effective_date="2026-10-01", program="core",
    )

    result = load_document(entry, tmp_path)

    assert "minimum equity injection" in result.text
    assert "Chapter 3" in result.text
    assert result.source_doc == "sop.docx"
    assert result.effective_date == "2026-10-01"
    assert result.program == "core"


def test_load_pdf_extracts_page_text(tmp_path):
    (tmp_path / "notice.pdf").write_bytes(b"%PDF-fake")
    entry = ManifestEntry(
        filename="notice.pdf", url="https://example.com/notice.pdf",
        file_type="pdf", effective_date="2025-08-28", program="7(a)",
    )
    fake_page = Mock()
    fake_page.extract_text.return_value = "SBA guaranty fees for FY2026 are reduced."
    fake_reader = Mock(pages=[fake_page])

    with patch("ingestion.loader.PdfReader", return_value=fake_reader):
        result = load_document(entry, tmp_path)

    assert "guaranty fees" in result.text
    assert result.source_doc == "notice.pdf"


def test_load_html_extracts_visible_text(tmp_path):
    (tmp_path / "cfr.html").write_text(
        "<html><body><p>A business is affiliated with another if...</p></body></html>",
        encoding="utf-8",
    )
    entry = ManifestEntry(
        filename="cfr.html", url="https://example.com/cfr.html",
        file_type="html", effective_date="2026-09-22", program="affiliation",
    )

    result = load_document(entry, tmp_path)

    assert "affiliated with another" in result.text


def test_load_document_raises_on_empty_text(tmp_path):
    (tmp_path / "empty.html").write_text("<html><body></body></html>", encoding="utf-8")
    entry = ManifestEntry(
        filename="empty.html", url="https://example.com/empty.html",
        file_type="html", effective_date="2026-01-01", program="core",
    )

    try:
        load_document(entry, tmp_path)
        raised = False
    except ValueError:
        raised = True
    assert raised


def test_load_document_raises_when_extraction_falls_below_min_chars(tmp_path):
    # A bot-check page or JS shell parses fine and yields a little text, so the
    # empty-text check alone would let it through and silently poison retrieval.
    (tmp_path / "blocked.html").write_text(
        "<html><body><p>Please enable JavaScript to view this page.</p></body></html>",
        encoding="utf-8",
    )
    entry = ManifestEntry(
        filename="blocked.html", url="https://example.com/cfr.html",
        file_type="html", effective_date="2026-01-01", program="affiliation",
        min_chars=50_000,
    )

    try:
        load_document(entry, tmp_path)
        raised = False
    except ValueError as exc:
        raised = "expected at least 50000" in str(exc)
    assert raised


def test_load_html_prefers_the_main_content_container(tmp_path):
    # eCFR wraps the regulation in div.part; the surrounding page is nav and
    # browser-support banners that chunk into near-duplicate noise.
    (tmp_path / "cfr.html").write_text(
        "<html><body>"
        "<div class='banner'>You are using an unsupported browser</div>"
        "<div class='part'>stub</div>"
        "<div class='part'>Affiliation is based on control over the concern.</div>"
        "</body></html>",
        encoding="utf-8",
    )
    entry = ManifestEntry(
        filename="cfr.html", url="https://example.com/cfr.html",
        file_type="html", effective_date="2026-01-01", program="affiliation",
    )

    result = load_document(entry, tmp_path)

    assert "Affiliation is based on control" in result.text
    assert "unsupported browser" not in result.text


def test_load_docx_builds_sections_from_heading_styles(tmp_path):
    # The SOP's structure lives in paragraph styles, not in the text itself.
    document = docx.Document()
    document.add_paragraph("Section A. Core Requirements", style="Heading 2")
    document.add_paragraph("Chapter 4: Fees", style="Heading 3")
    document.add_paragraph("Lenders must remit the guaranty fee.")
    document.add_paragraph("Chapter 5: Closing", style="Heading 3")
    document.add_paragraph("Closing requires SBA Form 155.")
    document.save(str(tmp_path / "sop.docx"))
    entry = ManifestEntry(
        filename="sop.docx", url="https://example.com/sop.docx",
        file_type="docx", effective_date="2026-10-01", program="core",
    )

    result = load_document(entry, tmp_path)

    bodies = {s.heading_path: s.text for s in result.sections}
    assert bodies[("Section A. Core Requirements", "Chapter 4: Fees")] == (
        "Lenders must remit the guaranty fee."
    )
    # Chapter 5 replaces Chapter 4 at the same level but stays under Section A.
    assert bodies[("Section A. Core Requirements", "Chapter 5: Closing")] == (
        "Closing requires SBA Form 155."
    )


def test_load_docx_skips_table_of_contents_paragraphs(tmp_path):
    # The TOC repeats every heading with a page number, which otherwise
    # retrieves as a near-duplicate of the real section.
    document = docx.Document()
    # The default template has no "toc 3"; the real SOP does, so add it.
    document.styles.add_style("toc 3", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("Chapter 4: Fees\t37", style="toc 3")
    document.add_paragraph("Chapter 4: Fees", style="Heading 3")
    document.add_paragraph("Lenders must remit the guaranty fee.")
    document.save(str(tmp_path / "sop.docx"))
    entry = ManifestEntry(
        filename="sop.docx", url="https://example.com/sop.docx",
        file_type="docx", effective_date="2026-10-01", program="core",
    )

    result = load_document(entry, tmp_path)

    assert "\t37" not in result.text


def test_load_html_builds_sections_from_heading_tags(tmp_path):
    (tmp_path / "cfr.html").write_text(
        "<html><body><div class='part'>"
        "<h2>Subpart A</h2>"
        "<h4>&#167; 121.103 How does SBA determine affiliation?</h4>"
        "<p>Concerns are affiliates when one controls the other.</p>"
        "</div></body></html>",
        encoding="utf-8",
    )
    entry = ManifestEntry(
        filename="cfr.html", url="https://example.com/cfr.html",
        file_type="html", effective_date="2026-09-22", program="affiliation",
    )

    result = load_document(entry, tmp_path)

    section = [s for s in result.sections if "affiliates" in s.text][0]
    assert section.heading_path == ("Subpart A", "§ 121.103 How does SBA determine affiliation?")


def test_load_pdf_strips_repeated_page_furniture(tmp_path):
    # Running headers/footers repeat on every page, landing mid-provision once
    # the pages are concatenated. Body text that happens to repeat must survive.
    (tmp_path / "notice.pdf").write_bytes(b"%PDF-fake")
    pages = []
    for number in (1, 2, 3):
        page = Mock()
        page.extract_text.return_value = "\n".join([
            "Federal Recycling Program Printed on Recycled Paper",
            "Upfront Fee for EWCP loans:",
            f"The fee on page {number} is 0.25% of the guaranteed portion.",
            "All loan amounts are based on the gross loan approval amount.",
            f"An additional fee may be due on page {number} when extended.",
            f"PAGE {number} of 3 EXPIRES: 10/1/26",
            "SBA Form 1353.3 (4-93) MS Word Edition; previous editions obsolete",
        ])
        pages.append(page)

    with patch("ingestion.loader.PdfReader", return_value=Mock(pages=pages)):
        result = load_document(entry_for_pdf(), tmp_path)

    assert "0.25%" in result.text
    assert "Federal Recycling Program" not in result.text
    assert "PAGE 1 of 3" not in result.text
    assert "SBA Form 1353.3" not in result.text
    # Repeats verbatim on every page but sits in the body, not the margin.
    assert "gross loan approval amount" in result.text


def test_load_pdf_detects_headings_without_markup(tmp_path):
    (tmp_path / "notice.pdf").write_bytes(b"%PDF-fake")
    page = Mock()
    page.extract_text.return_value = (
        "Upfront Fee for EWCP loans:\n"
        "For EWCP loans with a maturity of 12 months or less the fee is 0.25%.\n"
        "Additional Upfront Fee for 7(a) Loan Increases\n"
        "An additional fee is due on the increased amount.\n"
    )

    with patch("ingestion.loader.PdfReader", return_value=Mock(pages=[page])):
        result = load_document(entry_for_pdf(), tmp_path)

    paths = [s.heading_path for s in result.sections]
    assert ("Upfront Fee for EWCP loans:",) in paths
    assert ("Additional Upfront Fee for 7(a) Loan Increases",) in paths


def entry_for_pdf() -> ManifestEntry:
    return ManifestEntry(
        filename="notice.pdf", url="https://example.com/notice.pdf",
        file_type="pdf", effective_date="2025-08-28", program="7(a)",
    )


def _pdf_sections(tmp_path, page_text):
    (tmp_path / "notice.pdf").write_bytes(b"%PDF-fake")
    page = Mock()
    page.extract_text.return_value = page_text
    with patch("ingestion.loader.PdfReader", return_value=Mock(pages=[page])):
        return load_document(entry_for_pdf(), tmp_path).sections


def test_load_pdf_splits_a_heading_that_wraps_into_its_own_body(tmp_path):
    # A PDF wraps at a fixed width, so a short heading runs on into its body and
    # stops being a line ending in a colon. Left alone, the short-term fee rule
    # is filed under "exceeds 12 months" -- the opposite of what it answers.
    sections = _pdf_sections(tmp_path, (
        "For loans with a maturity that exceeds 12 months, the Upfront Fees are:\n"
        "Two percent of the guaranteed portion of the loan.\n"
        "For loans with a maturity of 12 months or less (Short-term loans): 0.25% of\n"
        "the guaranteed portion of the loan.\n"
    ))

    paths = [p for s in sections for p in s.heading_path]
    assert "For loans with a maturity of 12 months or less (Short-term loans):" in paths
    short_term = [s for s in sections if "12 months or less" in s.heading_path[0]][0]
    assert "0.25%" in short_term.text


def test_load_pdf_does_not_treat_the_tail_of_a_split_as_a_heading(tmp_path):
    # The tail begins mid-sentence on a capitalised word, so it would otherwise
    # satisfy the "short, mostly capitalised" heading test on its own.
    sections = _pdf_sections(tmp_path, (
        "Annual Service Fee for multiple 7(a) loans within 90 days: The Annual Service Fee\n"
        "is set for each loan on a standalone basis.\n"
    ))

    paths = [p for s in sections for p in s.heading_path]
    assert "Annual Service Fee for multiple 7(a) loans within 90 days:" in paths
    assert "The Annual Service Fee" not in paths


def test_load_pdf_does_not_split_fee_table_rows(tmp_path):
    # Each bullet is one row of a fee table; promoting rows to headings would
    # break the table into fragments that no longer read as a schedule.
    sections = _pdf_sections(tmp_path, (
        "For loans with a maturity that exceeds 12 months, the Upfront Fees are:\n"
        "• For loans of $150,000 or less: 2% of the guaranteed portion.\n"
        "• For loans of $150,001 to $700,000: 3% of the guaranteed portion.\n"
    ))

    assert len(sections) == 1
    assert "$150,000 or less: 2%" in sections[0].text
    assert "$150,001 to $700,000: 3%" in sections[0].text


def test_load_pdf_does_not_split_short_metadata_headers(tmp_path):
    # "CONTROL NO.: 5000-872051" is metadata, not a provision; splitting it
    # only produces a stub chunk.
    sections = _pdf_sections(tmp_path, (
        "CONTROL NO.: 5000-872051\n"
        "EFFECTIVE: August 28, 2025\n"
        "Each year SBA reviews certain fees payable by 7(a) Lenders.\n"
    ))

    paths = [p for s in sections for p in s.heading_path]
    assert "CONTROL NO.:" not in paths
